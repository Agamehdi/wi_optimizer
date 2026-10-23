"""Rumaila | Integrated Water Injection & Pattern Waterflood Optimizer.

Deploy to Streamlit Cloud with streamlit_app.py as entrypoint.
Use the sidebar navigation to access the ORIGINAL Gap and Pattern dashboards.
"""
from __future__ import annotations

import io
import re
from datetime import date
import numpy as np
import pandas as pd
import plotly.express as px
import plotly.graph_objects as go
import streamlit as st

import wi_gap as wg
import pattern_core as pc
import comparison_ui as cui
from optimizer_core import (OptimizerSettings, prepare_vow, well_scores,
                            approved_routes, optimize)

st.set_page_config(page_title='Rumaila | Water Injection Optimizer', page_icon='💧',layout='wide',initial_sidebar_state='expanded')
st.markdown('''<style>
.block-container{padding-top:1.4rem;max-width:1550px}
h1,h2,h3{letter-spacing:-.027em}
[data-testid="stMetric"]{background:#eff6fc;border:1px solid #cfdfeb;border-radius:13px;padding:14px}
[data-testid="stMetricLabel"]{font-weight:650;color:#244563}
.stTabs [data-baseweb="tab"]{font-weight:600}
</style>''',unsafe_allow_html=True)

NAVY='#16334f'; BLUE='#2367A1'; GREEN='#008b73'; ORANGE='#d48a35'; RED='#be4151'

def read_any(upload, sheet_name=0):
    b=upload.getvalue()
    if upload.name.lower().endswith('.csv'):
        return pd.read_csv(io.BytesIO(b),sep=None,engine='python',encoding='utf-8-sig')
    return pd.read_excel(io.BytesIO(b),sheet_name=sheet_name)


def format_view(df, numeric=None,mbd=True):
    out=df.copy()
    for c in numeric or []:
        if c in out:
            out[c]=pd.to_numeric(out[c],errors='coerce')/(1000 if mbd else 1)
    return out


def source_report(wells, facilities, pat, config, source_opts):
    selected=source_opts['date']
    d,unknown=wg.latest_wells(wells,selected,'Exclude and flag')
    if d.empty:raise ValueError('No BMT-assessed wells available at the selected date.')
    # Reservoir is optional in the well-source CSV; keep all downstream chart
    # and scoring tables consistent even when it was not supplied.
    if 'Reservoir' not in d.columns:
        d['Reservoir']=''
    tab,total,audit,unk=wg.calculate(wells,facilities,selected,
        missing='Exclude and flag',scenario=source_opts['scenario'],
        efficiency=False,gap_basis=source_opts['gap_basis'])
    cols=['Well','selected_injection']
    base_map=audit.loc[audit['capacity'].notna(),cols].set_index('Well')['selected_injection']
    d['baseline']=d.Well.map(base_map)
    current_candidates=[c for c in d.columns if re.sub(r'[^a-z0-9]','',c.lower()) in {
        'actualinjectionbbl', 'actualinjectionratebblday','currentinjectionbblday',
        'currentinjectionratebblday','actualinjectionbblday','currentinjratebblday'}]
    baseline_note='Modelled constrained injection (NOT measured actual injection)'
    if source_opts['use_actual'] and current_candidates:
        current=pd.to_numeric(d[current_candidates[0]].astype(str).str.replace(',','',regex=False),errors='coerce')
        if current.notna().all() and current.ge(0).all():
            d['baseline']=current
            baseline_note=f'Measured current injection supplied in {current_candidates[0]}'
        else:
            baseline_note+=' — optional current-rate column incomplete; reverted to modelled baseline'
    elif source_opts['use_actual']:
        baseline_note+=' — no matching measured Current Injection column; modelled baseline used'
    caps=wg.capacity_map(facilities,selected)
    return d,unknown,tab,total,audit,caps,baseline_note


def build_excel(data, gap_table, pattern, messages, settings, baseline_note):
    output=io.BytesIO()
    param=pd.DataFrame({'Parameter':['Mode','Additional supply (bbl/day)','Max increase (%)',
        'Max decrease (%)','Allowed new starts','Minimum VoW confidence','Demand ceiling',
        'Change penalty','Reroute penalty','Baseline','Solver status','Physical max status'],
        'Value':[settings.mode,settings.additional_supply_bpd,settings.max_increase_pct,
                 settings.max_decrease_pct,settings.allow_zero_start,settings.min_score_confidence,
                 settings.demand_limit,settings.change_penalty_vow,settings.route_penalty_vow,
                 baseline_note,data['status'],data['physical_max_status']]})
    notes=pd.DataFrame({'QC / notes':messages if messages else ['No additional warnings']})
    sheets={'Well_actions':data['wells'],'Manifold_actions':data['manifold'],
            'CPS_actions':data['cps'],'Gap_Analysis':gap_table,'Pattern_VoW':pattern,
            'Settings':param,'QC_Notes':notes}
    with pd.ExcelWriter(output,engine='openpyxl') as writer:
        for name,df in sheets.items():
            df.to_excel(writer,sheet_name=name,index=False)
            ws=writer.sheets[name];ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
            from openpyxl.styles import Font,PatternFill
            for cell in ws[1]:
                cell.fill=PatternFill('solid',fgColor='173A5D')
                cell.font=Font(name='Aptos',color='FFFFFF',bold=True)
            from openpyxl.utils import get_column_letter
            for col in ws.columns:
                head=col[0]
                length=max(len(str(x.value or '')) for x in list(col)[:150])
                ws.column_dimensions[get_column_letter(head.column)].width=min(44,max(13,length+2))
    return output.getvalue()


st.title('Water Injection | Integrated Optimization')
st.caption('Rumaila • Capacity gap × Pattern Value of Water × connectivity-constrained allocation')
st.info('**Engineering screening tool.** Results are proposed setpoint / reroute candidates for review, NOT approved operational instructions or a validated incremental-oil forecast. No cross-CPS transfers are permitted.')

with st.sidebar:
    st.header('01 · Source inputs')
    wf=st.file_uploader('Well RMT/BMT targets CSV',type=['csv'],key='optim_well')
    ff=st.file_uploader('Facility capacity CSV',type=['csv'],key='optim_fac')
    pf=st.file_uploader('Pattern Waterflood Excel / CSV',type=['csv','xlsx','xls'],key='optim_pattern')
    st.caption('All three main files are required. Original Gap/Pattern dashboards remain available in the Pages navigation.')
    pattern_sheet=0
    if pf and not pf.name.lower().endswith('.csv'):
        pattern_tabs=pd.ExcelFile(io.BytesIO(pf.getvalue())).sheet_names
        pattern_sheet=st.selectbox('Pattern workbook sheet',pattern_tabs,index=0)
    with st.expander('Optional network / well-pattern mapping'):
        routefile=st.file_uploader('Approved well → alternate manifold CSV',type=['csv','xlsx'],key='route_file',help='Columns Well, CPS, Allowed Manifold, Approved. Only Approved=Yes within the same CPS is considered.')
        wellpattern=st.file_uploader('Well → Pattern mapping CSV',type=['csv','xlsx'],key='well_pat',help='Columns Well, Pattern, optional Reservoir & Weight. Without this, manifold and CPS-level VoW scores are assigned as approximate proxies.')
        st.download_button('Approved routing CSV template',
            'Well,CPS,Allowed Manifold,Approved\n',file_name='approved_routes_template.csv',mime='text/csv')
        st.download_button('Well ↔ Pattern CSV template',
            'Well,Pattern,Reservoir,Weight\n',file_name='well_pattern_mapping_template.csv',mime='text/csv')
    st.header('02 · WI inputs')
    units=st.selectbox('Facility CAP units',['mbd','bbl/day'],index=0)
    date_fmt=st.selectbox('Well CSV date format',['Month/day/year','Day/month/year'])
    fac_fmt=st.selectbox('Facility date format',['Month/day/year','Day/month/year'])
    route_mode=st.selectbox('PWRI routing',['Separate PWRI facilities','SW only'])
    scenario=st.selectbox('Baseline target scenario',['RMT+OE','Agreed Target'])
    gap_basis=st.selectbox('Gap basis',['RMT+OE','RMT Target'])
    use_actual=st.checkbox('Use measured Current Injection, if complete',value=True,
        help='Requires every assessed well to have a numeric Current Injection [bbl/day] (or equivalent). Otherwise modelled constrained target is used and flagged.')
    st.header('03 · Pattern VoW')
    method=st.selectbox('Oil Supported method',
        [pc.BP_ANALYTICAL_METHOD,'Excel: Oil Supported','Excel: Support Fraction',
         'Produced Oil: user support %'])
    bo=pc.DEFAULT_BO.copy()
    if method==pc.BP_ANALYTICAL_METHOD:
        with st.expander('BP analytical fluid / Bo inputs'):
            for reservoir in pc.DEFAULT_BO:
                bo[reservoir]=st.number_input(f'{reservoir} Bo (rb/stb)',min_value=0.2,max_value=5.,value=float(bo[reservoir]),step=.01)
            bwi=st.number_input('Bwi (rb/stb)',value=1.,min_value=.1,step=.01)
            bwp=st.number_input('Bwp (rb/stb)',value=1.,min_value=.1,step=.01)
            dgor=st.number_input('Rp − Rs (scf/stb)',min_value=0.,value=0.,step=10.)
            bg=st.number_input('Bg (rb/scf)',min_value=0.,value=.005,format='%.5f')
    else:bwi,bwp,dgor,bg=1.,1.,0.,.005
    share=st.number_input('Produced-oil support proxy (%)',min_value=0.,max_value=100.,value=20.,step=5.) if method.startswith('Produced Oil') else 20.
    fallback=st.selectbox('VoW fallback scope',['Manifold + CPS','Manifold only'],
         help='Per-well Pattern score is most specific. CPS fallback is very approximate; set minimum confidence below to control whether CPS-only wells can change.')
    st.header('04 · Optimization')
    mode=st.selectbox('Water budget',['Conserve water','Additional supply'],index=0,
        help='Conserve: keep each CPS / PWRI water supply individually unchanged. Additional: allocate no more than each CPS baseline plus explicitly assigned extra supply; may reduce overloaded injection.')
    extra=st.number_input('Extra water available (mbd)',min_value=0.,value=0.,step=10.) if mode=='Additional supply' else 0.
    demand=st.selectbox('Upper well injection limit',['BMT capacity','RMT+OE','Agreed Target'],index=0)
    increase=st.slider('Max increase per well (%)',0,500,200,10)
    decrease=st.slider('Max decrease per well (%)',0,100,100,10)
    start_zero=st.checkbox('Allow currently zero-rate injectors to start (review required)',value=False)
    confidence=st.select_slider('Minimum VoW matching confidence',options=[0.,.35,.55,.75,.90,1.],value=.75,
        help='1.0=direct Well/Pattern; .90=Manifold+Reservoir; .75=Manifold; .55=CPS+Reservoir; .35=CPS-only. Unscored/low-confidence rates are kept as close as capacity allows.')
    change_cost=st.number_input('Rate-change penalty (VoW units)',0.,1.,.002,step=.001,format='%.3f')
    route_cost=st.number_input('Manifold reroute penalty (VoW units)',0.,1.,.01,step=.005,format='%.3f')
    timeout=st.slider('Optimizer time limit (sec)',5,120,45,5)

if not (wf and ff and pf):
    a,b,c=st.columns(3)
    a.metric('Water injection input','Required' if not wf else 'Uploaded')
    b.metric('Facility capacity','Required' if not ff else 'Uploaded')
    c.metric('Pattern analytics','Required' if not pf else 'Uploaded')
    st.subheader('How it makes recommendations')
    st.markdown('''The optimizer reads the **exact well and CPS/manifold capacity constraints** from Gap Analysis and the **raw, non-normalized VoW coefficients** from Pattern Waterflood. It can adjust existing well rates within their network, find unused capacity in higher-value manifolds, and optionally propose approved same-CPS reroutes. A well always keeps its original Pattern VoW score when rerouted. Without an explicit route file, **no manifold changes** are permitted.

Upload the three existing files in the sidebar; add connectivity/Pattern mapping if available to strengthen the recommendations.''')
    st.stop()

try:
    raw_wells=wg.read_csv(wf)
    wells=wg.prepare_wells(raw_wells,dayfirst=date_fmt=='Day/month/year',
        route='PWRI priority' if route_mode.startswith('Separate') else 'SW')
    facilities=wg.prepare_facilities(wg.read_csv(ff),units=units,dayfirst=fac_fmt=='Day/month/year')
    pattern_raw=read_any(pf,sheet_name=pattern_sheet)
    date_default=wells.date.max().date()
    analysis_date=st.date_input('Analysis snapshot (well & capacity effective date)',date_default,
            min_value=wells.date.min().date(),max_value=date.today().replace(year=date.today().year+5))
    if pd.Timestamp(analysis_date)<wells.date.min():raise ValueError('No wells exist at this date.')
    opts={'date':analysis_date,'scenario':scenario,'gap_basis':gap_basis,'use_actual':use_actual}
    d,missing,gap_table,gap_total,audit,caps,baseline_note=source_report(wells,facilities,pattern_raw,None,opts)
    pattern,issues=prepare_vow(pattern_raw,method,bo,bwi,bwp,dgor,bg,share)
    if pattern.empty:raise ValueError('No Pattern rows have valid injected volume and supported-oil coefficient.')
    m=read_any(wellpattern) if wellpattern else None
    scored,match_issues=well_scores(d,pattern,m,fallback=fallback)
    routes,route_issues=approved_routes(scored,caps,read_any(routefile) if routefile else None)
    extra_cps=(st.selectbox('CPS / PWRI facility receiving additional water',sorted(scored.facility.unique()),help='Explicit extra water is assigned to this facility only. No cross-CPS borrowing is assumed.') if mode=='Additional supply' and extra>0 else None)
    config=OptimizerSettings(mode=mode,additional_supply_bpd=extra*1000,extra_by_facility={extra_cps:extra*1000} if extra_cps else {},demand_limit=demand,
        max_increase_pct=increase,max_decrease_pct=decrease,allow_zero_start=start_zero,
        min_score_confidence=confidence,change_penalty_vow=change_cost,route_penalty_vow=route_cost,
        solver_time_limit_sec=timeout)
    warnings=issues+match_issues+route_issues
    if len(missing):warnings.insert(0,f'{len(missing)} wells have missing BMT capacity: excluded from optimization, NOT assessed as free injection capacity.')
    if scored['VoW raw'].isna().any():
        warnings.append('Some wells have no matched Pattern VoW; their rates are protected against discretionary reallocation.')
    st.caption(f'**Baseline:** {baseline_note}  |  **Snapshot:** {analysis_date}  |  **PWRI:** {route_mode}')
except Exception as e:
    st.error(f'Input QC failed: {e}')
    st.stop()

base_total=float(scored.baseline.sum())
col1,col2,col3,col4=st.columns(4)
col1.metric('Wells assessed',f'{len(scored):,}')
col2.metric('Baseline injection',f'{base_total/1000:,.1f} mbd')
col3.metric('Scored wells',f'{int((scored["Score confidence"]>=confidence).sum()):,}')
col4.metric('Approved alternate edges',sum(max(0,len(v)-1) for v in routes.values()))

with st.expander('Pre-optimization input QC and source matching',expanded=False):
    st.dataframe(scored[['Well','Reservoir','facility','manifold','baseline','capacity','VoW raw','VoW source','Score confidence']].sort_values(['facility','manifold','Well']),hide_index=True,use_container_width=True)
    if warnings:
        for note in warnings[:40]:st.warning(note)
    st.dataframe(pattern[['Pattern','Reservoir','CPS','Manifold','Injection','Supported','raw_vow']].head(150),hide_index=True,use_container_width=True)

if 'optimizer_bundle' not in st.session_state:st.session_state.optimizer_bundle=None
if st.button('⚡ Run connected-network optimization',type='primary',use_container_width=True):
    try:
        with st.spinner('Solving physical well / manifold / CPS constraints...'):
            result=optimize(scored,caps,routes,config)
        st.session_state.optimizer_bundle={'data':result,'gap':gap_table,'pattern':pattern,
            'notes':warnings,'config':config,'baseline_note':baseline_note,
            'identity':(wf.getvalue(),ff.getvalue(),pf.getvalue(),
                routefile.getvalue() if routefile else None,
                wellpattern.getvalue() if wellpattern else None,
                str(analysis_date),str(config),str(pattern_sheet),method,str(bo),fallback,scenario,route_mode,confidence,share,bwi,bwp,dgor,bg)}
    except ValueError as exc:
        st.error(f'No feasible solution under these inputs: {exc}')
        st.info('If current injection already exceeds capacity, try **Additional supply** with 0 extra water. This permits necessary reductions rather than forcing exact conservation. Review well-rate minimums and source capacity QC.')

bundle=st.session_state.optimizer_bundle
identity=(wf.getvalue(),ff.getvalue(),pf.getvalue(),routefile.getvalue() if routefile else None,
          wellpattern.getvalue() if wellpattern else None,str(analysis_date),str(config),str(pattern_sheet),method,
          str(bo),fallback,scenario,route_mode,confidence,share,bwi,bwp,dgor,bg)
if bundle and bundle['identity']!=identity:
    st.warning('Settings or uploaded data changed. Run optimization again for updated results.')
    bundle=None

if bundle:
    result=bundle['data']; rec=result['wells'];mf=result['manifold'];fac=result['cps']
    st.subheader('Optimization outcome')
    st.caption(f"Solver: **{result['status']}**. Physical-max solve: {result['physical_max_status']}.")
    k1,k2,k3,k4,k5=st.columns(5)
    k1.metric('Baseline',f'{rec.baseline.sum()/1000:,.2f} mbd')
    k2.metric('Optimized',f'{rec["Optimized injection"].sum()/1000:,.2f} mbd',delta=f'{rec["Delta injection"].sum()/1000:+,.2f} mbd')
    k3.metric('Reachable injection max',f'{result["physical_max"]/1000:,.2f} mbd' if np.isfinite(result['physical_max']) else 'n/a')
    k4.metric('Suggested well changes',f'{int(rec.Action.ne("Maintain").sum())}')
    k5.metric('Approved manifold reroutes',f'{int(rec["Recommended manifold"].ne(rec.manifold).sum())}')
    oil_delta=rec['Oil-supported proxy change (stb/d)'].sum()
    st.info(f'**Oil-supported proxy change:** {oil_delta:+,.0f} stb/day under fixed VoW assumptions. This is a planning index, NOT a causal incremental-oil forecast. Validate with reservoir simulation, pressure/VRR, injectivity and facilities before implementation.')
    if (fac['Before ullage']<-1).any() or (mf['Before ullage']<-1).any():
        st.warning('The supplied baseline exceeds at least one manifold/CPS limit. Negative BEFORE ullage is a QC/operational constraint breach requiring investigation; optimization uses supplied capacity limits.')

    overview, compare_tab, actions, audit_tab=st.tabs(['CPS / Manifold overview','📊 Original vs Optimized','Recommended actions','Engineering QC & download'])
    with overview:
        plot_cps=fac.melt(id_vars='CPS',value_vars=['Current injection','Optimized injection','Capacity'],var_name='Scenario',value_name='Rate')
        plot_cps['Rate (mbd)']=plot_cps.Rate/1000
        fig=px.bar(plot_cps,x='CPS',y='Rate (mbd)',color='Scenario',barmode='group',
            color_discrete_map={'Current injection':BLUE,'Optimized injection':GREEN,'Capacity':ORANGE},title='CPS | baseline vs optimized vs maximum network capacity')
        fig.update_layout(template='plotly_white',height=430,title_font_color=NAVY,legend=dict(orientation='h',y=-.25))
        st.plotly_chart(fig,use_container_width=True)
        st.markdown('**CPS summary (mbd)**')
        st.dataframe(format_view(fac,['Current injection','Optimized injection','Capacity','Before ullage','After ullage','Change']),hide_index=True,use_container_width=True)
        mf_plot=mf.copy();mf_plot['Label']=mf_plot.CPS+' / '+mf_plot.Manifold
        mf_plot=mf_plot.sort_values('Change',ascending=False)
        if len(mf_plot):
            fig2=go.Figure()
            fig2.add_bar(y=mf_plot.Label,x=mf_plot['Current injection']/1000,name='Baseline',orientation='h',marker_color=BLUE)
            fig2.add_bar(y=mf_plot.Label,x=mf_plot['Optimized injection']/1000,name='Optimized',orientation='h',marker_color=GREEN)
            fig2.update_layout(barmode='group',height=max(380,35*len(mf_plot)),template='plotly_white',title='Manifold injection reallocation',xaxis_title='Injection (mbd)',yaxis={'autorange':'reversed'},margin={'l':130,'b':60,'t':65})
            st.plotly_chart(fig2,use_container_width=True)
        st.markdown('**Manifold changes and remaining ullage (mbd)**')
        st.dataframe(format_view(mf,['Current injection','Optimized injection','Capacity','Before ullage','After ullage','Change']),hide_index=True,use_container_width=True)
    with compare_tab:
        cui.render_comparison(result,caps,bundle['gap'],bundle['pattern'],mapping=m,
              target_default='RMT Target' if gap_basis=='RMT Target' else 'RMT+OE Target',
              baseline_note=baseline_note,analysis_date=analysis_date)
    with actions:
        changed=rec.loc[rec.Action.ne('Maintain')].copy()
        changed['Abs delta']=changed['Delta injection'].abs()
        changed=changed.sort_values('Abs delta',ascending=False)
        if changed.empty:st.success('No discretionary well changes identified under the selected constraints.')
        else:
            for _,item in changed.head(12).iterrows():
                action=item['Action']; old=item['manifold']; new=item['Recommended manifold']
                st.markdown(f"**{item['Well']} · {item['facility']}** — {action} | {old} → {new} | **{item['Delta injection']/1000:+.2f} mbd** | VoW {item['VoW raw']:.3f} ({item['VoW source']})")
        export_cols=['Well','Reservoir','facility','manifold','Recommended manifold','baseline','Optimized injection','Delta injection','capacity','physical_upper','VoW raw','VoW source','Score confidence','Oil-supported proxy change (stb/d)','Action']
        st.dataframe(format_view(changed[export_cols],['baseline','Optimized injection','Delta injection','capacity','physical_upper']),hide_index=True,use_container_width=True)
        reroute=changed[changed['Recommended manifold'].ne(changed.manifold)]
        if not reroute.empty:
            st.warning('Reroutes are candidate engineering changes only. Confirm headers, hydraulics, isolation/valves, tie-in availability, pressure rating and installation costs before considering execution.')
        st.download_button('Download recommended actions CSV',changed[export_cols].to_csv(index=False).encode('utf-8-sig'),file_name='WI_Optimization_Action_List.csv',mime='text/csv')
    with audit_tab:
        st.markdown('**Key engineering assumptions**')
        st.markdown('''- A well can only route to its **original manifold**, unless an **Approved=Yes** route exists within the **same CPS**, with a dated target-manifold capacity. PWRI branches are always locked separately.
- **No implicit cross-CPS transfer** of wells OR allocated injection water: each CPS / PWRI has its own budget; no multiple simultaneous manifold assignments, no invented physical tie-ins. Each well's score is tied to its original pattern; new manifold does **not** change its VoW.
- Constraints are: well BMT limit (plus optional chosen target cap), per-well increase/decrease, physical manifold and CPS maximum, water supply, and binary routing. Solver selects total supported-oil index subject to those constraints.
- The source Pattern data's **raw VoW** is used, not display normalization. VoW may not represent the **marginal** response to increasing water. Apply valid pressure, injection/sweep, integrity and water-availability checks independently.
- Unscored or insufficiently confident wells are held at current rate where capacity allows. Current-rate column is preferred only when complete; otherwise constrained *modelled* target is baseline.
- Before field execution, verify well injectivity/injection pressure, pressure gradients, voidage replacement, faults/compartments, pressure ceiling, channeling, manifold operating envelope and offset producers.''')
        if bundle['notes']:
            st.markdown('**QC / rejected connectivity edges**')
            for note in bundle['notes'][:80]:st.caption('• '+note)
        report=build_excel(result,bundle['gap'],bundle['pattern'],bundle['notes'],bundle['config'],bundle['baseline_note'])
        st.download_button('⬇ Download complete Excel: action plan + CPS/Manifold + QC',report,file_name='Rumaila_Integrated_WI_Optimization.xlsx',mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',type='primary',use_container_width=True)
