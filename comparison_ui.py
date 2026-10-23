"""Interactive engineering-grade Original vs Optimized dashboard in Streamlit."""
from __future__ import annotations
import io
import pandas as pd
import numpy as np
import streamlit as st
import plotly.io as pio

import comparison_views as cv
import gap_summary as gs

RATE_NAMES=['RMT Target','RMT+OE Target','Agreed Target','BMT Well Capacity',
            'MF Capacity','CPS Capacity','Injection','Delivery Gap',
            'MF Target Shortfall','CPS Target Shortfall','MF Ullage','CPS Ullage',
            'Injection Change','Well Physical Shortfall','Target Surplus']


def mbd_view(df):
    d=df.copy()
    for col in d.columns:
        if any(str(col).startswith(n) for n in RATE_NAMES) and 'BMT wells' not in str(col):
            d[col]=pd.to_numeric(d[col],errors='coerce')/1000
    return d


def _show_decision_table(d):
    fmt={c:'{:+,.2f}' if 'Change' in c else '{:,.2f}' for c in d.select_dtypes(include='number')
         if c not in ('Wells Before','Wells After','Missing BMT wells')}
    st.dataframe(d.style.format(fmt,na_rep='—'),use_container_width=True,hide_index=True,height=min(680,100+len(d)*35))


def _download_html(figures):
    blocks=[]
    for i,(title,fig) in enumerate(figures):
        blocks.append(f'<h2>{title}</h2>'+pio.to_html(fig,full_html=False,include_plotlyjs=True if i==0 else False))
    return ('<!doctype html><html><head><meta charset="utf-8"><title>Rumaila WI Optimization | Before / After</title>'
            '<style>body{font:14px Arial,sans-serif;max-width:1450px;margin:30px auto;color:#173553}'
            'h2{margin:46px 20px 8px}header{margin:20px}</style></head><body>'
            '<header><h1>Rumaila | Water Injection — Original vs Optimized</h1>'
            '<p>Engineering screening • Oil Supported is a fixed-VoW proxy, not forecast incremental oil.</p></header>'
            + ''.join(blocks) +'</body></html>').encode('utf-8')


def render_comparison(result,caps,source_gap,pattern,mapping=None,target_default='RMT+OE Target',baseline_note='',analysis_date=None):
    """Called only AFTER optimizer result exists and has passed input QC."""
    st.header('Original vs Optimized | Performance & Capacity Gap')
    st.caption('All bar charts are calculated from the optimizer’s dated well-by-well result, with the same CPS and manifold capacities in both scenarios.')

    s1,s2,s3=st.columns([1.2,1.1,1])
    with s1:
        target=st.selectbox('Target basis', ['RMT+OE Target','RMT Target','Agreed Target'],
                 index=['RMT+OE Target','RMT Target','Agreed Target'].index(target_default),
                 key='cmp_target')
    with s2:
        mask_changed=st.checkbox('Show changed points only (scatter)',value=False,key='cmp_changed')
    with s3:
        top=int(st.selectbox('Manifold chart scope', [10,15,25,1000],index=2,
                             format_func=lambda x: 'All manifolds' if x==1000 else f'Top {x} changes',key='cmp_mf_limit'))

    try:
        comp=cv.prepare_comparisons(result,caps,source_gap,pattern,mapping,target)
    except (ValueError,KeyError,TypeError) as e:
        st.error('Cannot compute before/after comparisons: '+str(e))
        return None
    metrics=comp['metrics']
    cols=st.columns(5)
    cols[0].metric('Injection | Original',f"{metrics['injection_before_mbd']:,.1f} mbd")
    cols[1].metric('Injection | Optimized',f"{metrics['injection_after_mbd']:,.1f} mbd",
                   delta=f"{metrics['injection_after_mbd']-metrics['injection_before_mbd']:+,.1f} mbd")
    cols[2].metric('Delivery gap | Original',f"{metrics['cps_delivery_gap_before_mbd']:,.1f} mbd")
    cols[3].metric('Delivery gap | Optimized',f"{metrics['cps_delivery_gap_after_mbd']:,.1f} mbd",
                   delta=f"{metrics['cps_delivery_gap_after_mbd']-metrics['cps_delivery_gap_before_mbd']:+,.1f} mbd",delta_color='inverse')
    cols[4].metric('Oil-supported PROXY Δ',f"{metrics['support_after_mbd']-metrics['support_before_mbd']:+,.2f} mbd",
                   help='Fixed well/pattern VoW coefficient × change in allocated water. NOT predicted incremental oil.')
    # Inspect the computed source comparison, not the display variable `mf`
    # (which is only created below). This previously raised UnboundLocalError
    # as soon as comparison calculations succeeded.
    source_mf=comp['capacity_manifold']
    if ('Missing BMT wells' in source_mf.columns) and source_mf['Missing BMT wells'].sum()>0:
        st.warning('The source Gap Analysis includes wells without assessed BMT capacity. Optimization and new target-delivery comparisons include assessed wells only. '
                   'Review the Missing BMT column and original source table before treating this as a field-wide gap.')
    st.info(f"**Basis:** {baseline_note}. Supported-oil values are a **fixed-VoW engineering screening proxy**, not field production forecasts. "
            f"VoW score coverage: Original {metrics['score_coverage_before']:.0f}% / Optimized {metrics['score_coverage_after']:.0f}% of injection. "
            "Unmatched wells are not assigned fabricated supported-oil volumes.")

    cps=mbd_view(comp['capacity_cps'])
    mf=mbd_view(comp['capacity_manifold'])
    fac_names=cps.CPS.astype(str).tolist()
    selection=st.multiselect('Filter CPS / PWRI facilities',fac_names,default=fac_names,key='cmp_cps_filter')
    if not selection:
        st.warning('Select one or more CPS / PWRI facilities.')
        return comp
    cps=cps[cps.CPS.isin(selection)].copy()
    mf=mf[mf.CPS.isin(selection)].copy()
    cs=comp['scatter_cps'];cs=cs[cs.CPS.isin(selection)].copy()
    ms=comp['scatter_manifold'];ms=ms[ms.CPS.isin(selection)].copy()
    ps=comp['scatter_pattern'];ps=ps[ps.CPS.isin(selection)].copy() if not ps.empty else ps
    st.caption('**Blue = Original • Green = Optimized.** Negative ullage in Original indicates a supplied current-rate/capacity conflict; after values must respect the supplied hard limits.')
    charts=[]
    a,b,c,d=st.tabs(['Capacity Gap & Tables','VoW × Oil Supported','Ullage by CPS / Manifold','Optimization Decision Summary'])
    with a:
        st.markdown('### Original vs Optimized — Full Capacity Gap Summary')
        st.caption('Exact 18-column layout of the supplied WI Gap Analysis source screenshot. Original is the selected dated baseline; '
                   'Optimized is the exact solved well allocation, including approved same-CPS manifold changes. '
                   'All values are **mbd**. Scroll horizontally to view all 18 columns. The separate Changes tab retains the decision-impact columns.')
        original_summary, optimized_summary, gap_delta = gs.build_gap_summary(comp, target)
        original_summary = original_summary[original_summary['CPS'].isin(selection + ['TOTAL'])].copy()
        optimized_summary = optimized_summary[optimized_summary['CPS'].isin(selection + ['TOTAL'])].copy()
        gap_delta = gap_delta[gap_delta['CPS'].isin(selection)].copy()
        # Selection-aware TOTAL rows: never present full-field totals beneath a
        # filtered CPS selection. Sum facility-only measures only once per CPS.
        for table in (original_summary, optimized_summary):
            grouped=table[table['CPS'].ne('TOTAL')]
            tidx=table.index[table['CPS'].eq('TOTAL')]
            if len(tidx):
                for field in gs.RATE_FIELDS:
                    v=grouped.drop_duplicates('CPS')[field] if field in gs.FACILITY_FIELDS else grouped[field]
                    table.loc[tidx,field]=v.sum() if v.notna().all() else np.nan
                table.loc[tidx,'Missing BMT wells']=grouped['Missing BMT wells'].sum()
        original_total=original_summary[original_summary.CPS.eq('TOTAL')].iloc[0]
        optimized_total=optimized_summary[optimized_summary.CPS.eq('TOTAL')].iloc[0]
        j1,j2,j3=st.columns(3)
        j1.metric('Injection improvement',
            f"{optimized_total['Manifold Injection Rate']-original_total['Manifold Injection Rate']:+,.1f} mbd")
        j2.metric('MF delivery gap reduction',
            f"{original_total['Delivery Gap to Selected Target']-optimized_total['Delivery Gap to Selected Target']:+,.1f} mbd")
        j3.metric('CPS unused capacity change',
            f"{optimized_total['CPS Ullage']-original_total['CPS Ullage']:+,.1f} mbd")
        summary_tabs=st.tabs(['Original — Before','Optimized — After','Changes by Manifold'])
        with summary_tabs[0]:
            st.markdown(gs.gap_table_html(original_summary,'ORIGINAL',target,analysis_date),unsafe_allow_html=True)
        with summary_tabs[1]:
            st.markdown(gs.gap_table_html(optimized_summary,'OPTIMIZED',target,analysis_date),unsafe_allow_html=True)
        with summary_tabs[2]:
            st.caption('Positive delivery-gap reduction is favorable. Negative manifold ullage change means '
                       'previously unused capacity is being utilized. Physical capacity shortfall remains unchanged '
                       'without rerouting or hardware changes.')
            gap_delta = gap_delta.sort_values('Injection Δ (mbd)',key=lambda x:x.abs(),ascending=False)
            _show_decision_table(gap_delta)
        if original_total['Missing BMT wells']:
            st.warning('Some BMT well capacities are not assessed. The well-constraint column is n.a. '
                       'for affected manifolds and cannot be interpreted as a complete physical gap.')
        st.caption('IIC and INJ_RATE @ OE in this comparison represent the SAME allocated CPS injection '
                   'in the selected Original/Optimized scenario. These are not separate RMT+OE recomputations. '
                   'Well and facility capacities are unchanged; only approved well setpoints/routes may change.')
        st.divider()
        st.markdown('### 01 · Target delivery shortfall')
        fig=cv.bar_before_after(cps,'CPS','Delivery Gap Before','Delivery Gap After',
            f'{target} | Remaining delivery gap by CPS','Shortfall to target (mbd)')
        st.plotly_chart(fig,use_container_width=True,key='gap_cps');charts.append(('CPS target-delivery gap',fig))
        mf_view=mf.assign(Label=mf.CPS+' / '+mf.Manifold)
        fig2=cv.bar_before_after(mf_view,'Label','Delivery Gap Before','Delivery Gap After',
              'Target delivery gap by manifold','Shortfall (mbd)',horizontal=True,
              top=None if top==1000 else top)
        st.plotly_chart(fig2,use_container_width=True,key='gap_mf');charts.append(('Manifold target-delivery gap',fig2))
        st.markdown('### 02 · Physical gap vs allocation gap')
        st.caption('**Physical shortfall** compares RMT/Agreed demand with facility maximum; it does not disappear when injection is redistributed. '
                   '**Delivery gap** compares actual/modelled injection against the selected target. '
                   'With approved reroutes, targets are attributed to the proposed destination manifold in AFTER.')
        gap_ch=cv.bar_before_after(cps,'CPS','CPS Target Shortfall Before','CPS Target Shortfall After',
                'Structural CPS capacity shortfall (should not change with rate allocation)',
                'Demand above CPS limit (mbd)')
        st.plotly_chart(gap_ch,use_container_width=True,key='structural_gap');charts.append(('Structural capacity shortfall',gap_ch))
        with st.expander('Original Gap Analysis export (source calculations)',expanded=False):
            st.dataframe(source_gap,use_container_width=True,hide_index=True)
        st.markdown('**Before/After Capacity Gap Table — CPS (mbd)**')
        ccols=['CPS',f'{target} Before','CPS Capacity','Injection Before','Injection After',
               'Delivery Gap Before','Delivery Gap After','CPS Target Shortfall Before',
               'CPS Target Shortfall After','CPS Ullage Before','CPS Ullage After','Injection Change']
        _show_decision_table(cps[ccols])
        st.markdown('**Before/After Capacity Gap Table — Manifold (mbd)**')
        mcols=['CPS','Manifold',f'{target} Before',f'{target} After',
               'MF Capacity','BMT Well Capacity Before','BMT Well Capacity After',
               'Injection Before','Injection After', 'Delivery Gap Before','Delivery Gap After',
               'MF Target Shortfall Before','MF Target Shortfall After',
               'Well Physical Shortfall Before','Well Physical Shortfall After',
               'MF Ullage Before','MF Ullage After','Injection Change']
        if 'Missing BMT wells' in mf.columns:mcols.append('Missing BMT wells')
        _show_decision_table(mf[mcols])

    with b:
        st.markdown('### 03 · Water injection vs oil supported — Original and optimized')
        st.caption('Same pattern-based VoW coefficient is applied to both well allocations. Lines connect a single CPS/manifold Original→Optimized state. '
                   'A move on these charts reflects a proposed change in injection and **proxy** supported oil, not proof of new reservoir response.')
        cvs=cv.injection_vs_support(cs,['CPS'],'CPS | Injection vs Oil-supported proxy',mask_changed)
        st.plotly_chart(cvs,use_container_width=True,key='supp_cps');charts.append(('CPS injection versus supported',cvs))
        mf_fig=cv.injection_vs_support(ms,['CPS','Manifold'], 'CPS × Manifold | Injection vs Oil-supported proxy',mask_changed)
        st.plotly_chart(mf_fig,use_container_width=True,key='supp_mf');charts.append(('Manifold injection versus supported',mf_fig))
        st.markdown('### 04 · VoW vs Oil Supported')
        vv=cv.vow_vs_supported(cs,['CPS'],'CPS | Weighted VoW vs supported-oil proxy',mask_changed)
        st.plotly_chart(vv,use_container_width=True,key='vow_cps');charts.append(('VoW vs supported CPS',vv))
        vv2=cv.vow_vs_supported(ms,['CPS','Manifold'],'Manifold | Weighted VoW vs supported-oil proxy',mask_changed)
        st.plotly_chart(vv2,use_container_width=True,key='vow_mf');charts.append(('VoW vs supported manifold',vv2))
        st.markdown('### 05 · Pattern level')
        if not ps.empty:
            st.success('Well→Pattern mapping detected. Pattern rates are reconstructed from the SAME dated well baseline/optimized snapshot.')
            pf=cv.injection_vs_support(ps,['CPS','Pattern'],'Pattern | Injection vs Oil-supported proxy',mask_changed)
            st.plotly_chart(pf,use_container_width=True,key='supp_pattern');charts.append(('Pattern injection versus supported',pf))
            vp=cv.vow_vs_supported(ps,['CPS','Pattern'],'Pattern | Weighted VoW vs Oil-supported proxy',mask_changed)
            st.plotly_chart(vp,use_container_width=True,key='vow_pattern');charts.append(('VoW vs supported pattern',vp))
            st.dataframe(ps,use_container_width=True,hide_index=True)
        else:
            st.warning('A defensible Pattern AFTER plot requires a **Well → Pattern mapping CSV**. Upload this in the sidebar. '
                       'The existing source Pattern scatter below is only the historical uploaded data (NOT an optimized forecast).')
            ph=cv.source_pattern_reference(pattern)
            st.plotly_chart(ph,use_container_width=True,key='pattern_reference');charts.append(('Original Pattern input reference',ph))
        for msg in comp['pattern_issues'][:12]:st.caption('QC: '+msg)

    with c:
        st.markdown('### 06 · Available ullage before vs after')
        u1=cv.bar_before_after(cps,'CPS','CPS Ullage Before','CPS Ullage After',
                    'CPS | Total capacity minus allocated injection','Remaining CPS ullage (mbd)')
        st.plotly_chart(u1,use_container_width=True,key='ullage_cps');charts.append(('CPS ullage',u1))
        u2=cv.bar_before_after(mf.assign(Label=mf.CPS+' / '+mf.Manifold),'Label',
                       'MF Ullage Before','MF Ullage After','Manifold | Remaining capacity',
                       'Remaining manifold ullage (mbd)',horizontal=True,
                       top=None if top==1000 else top)
        st.plotly_chart(u2,use_container_width=True,key='ullage_mf');charts.append(('Manifold ullage',u2))
        st.caption('Ullage = capacity – injection. A reduction in ullage means more capacity is used; it is not automatically a benefit unless the injection is technically justified and water is available.')
        st.dataframe(mf[['CPS','Manifold','MF Capacity','Injection Before','Injection After',
                         'MF Ullage Before','MF Ullage After','Injection Change']].sort_values('Injection Change',ascending=False),
                         use_container_width=True,hide_index=True)

    with d:
        st.markdown('### 07 · Allocation impact: which changes drive the result?')
        w=result['wells'].copy()
        w=w[w.facility.isin(selection)].copy()
        if 'Reservoir' not in w.columns:
            w['Reservoir']=''
        if 'VoW source' not in w.columns:
            w['VoW source']='Unavailable'
        w['Original (mbd)']=w.baseline/1000
        w['Optimized (mbd)']=w['Optimized injection']/1000
        w['Change (mbd)']=w['Delta injection']/1000
        w['Supported proxy Δ (mbd)']=w['Oil-supported proxy change (stb/d)']/1000
        w['Network action']=np.where(w['Recommended manifold'].ne(w.manifold),'Approved manifold reroute',
                             np.where(w['Change (mbd)']>.001,'Increase setpoint',
                                      np.where(w['Change (mbd)']<-.001,'Decrease setpoint','Maintain')))
        w=w.sort_values('Supported proxy Δ (mbd)',ascending=False)
        st.dataframe(w[['Well','Reservoir','facility','manifold','Recommended manifold',
                        'Original (mbd)','Optimized (mbd)','Change (mbd)','VoW raw',
                        'Score confidence','VoW source','Supported proxy Δ (mbd)','Network action']],
             hide_index=True,use_container_width=True)
        bsup=metrics['support_before_mbd'];asup=metrics['support_after_mbd']
        s1,s2,s3=st.columns(3)
        s1.metric('Supported proxy | Original',f'{bsup:,.2f} mbd')
        s2.metric('Supported proxy | Optimized',f'{asup:,.2f} mbd')
        s3.metric('Allocation proxy improvement',f'{asup-bsup:+,.2f} mbd')
        st.markdown('**Decision criteria / engineering limitations**')
        st.markdown('''- All Original and Optimized values are taken from the **same dated well snapshot** and **same optimizer solution**.
- CPS / PWRI water budgets remain separate, except for additional water **explicitly assigned to a facility** in the sidebar.
- Manifold rerouting requires **approved exact well-to-manifold connectivity** within the original CPS. A different manifold is **not** a higher VoW reservoir; the well score stays with the well/pattern.
- If current injection is modelled RMT/BMT-constrained rate instead of measured actuals, plots are **modelled Original vs Optimized**, not measured performance.
- Water supply, pressures, well injectivity, reservoir sweep response, water quality and operational approvals still need validation. Normalized display VoW is never an optimization input.
- Without measured incremental response / saturation and pressure forecasts, **Oil Supported is a linear screening proxy**, not additional future oil production.''')

    excel=io.BytesIO()
    with pd.ExcelWriter(excel,engine='openpyxl') as writer:
        # Added snapshot-style before/after and directly comparable delta exports.
        full_orig, full_opt, full_delta = gs.build_gap_summary(comp, target)
        for sh, dataset in [('Gap_Summary_Original',full_orig),('Gap_Summary_Optimized',full_opt),('Gap_Manifold_Changes',full_delta)]:
            dataset.to_excel(writer,sheet_name=sh,index=False)
        mbd_view(comp['capacity_cps']).to_excel(writer,sheet_name='CPS_Before_After',index=False)
        mbd_view(comp['capacity_manifold']).to_excel(writer,sheet_name='MF_Before_After',index=False)
        comp['scatter_cps'].to_excel(writer,sheet_name='CPS_Support_Scatter',index=False)
        comp['scatter_manifold'].to_excel(writer,sheet_name='MF_Support_Scatter',index=False)
        if not comp['scatter_pattern'].empty:
            comp['scatter_pattern'].to_excel(writer,sheet_name='Pattern_Support_Scatter',index=False)
        result['wells'].to_excel(writer,sheet_name='Optimizer_Well_Actions',index=False)
        meta=pd.DataFrame({'Setting':['Target','Baseline','Oil supported definition','Raw VoW','Pattern Plot'],
         'Value':[target,baseline_note,'Fixed pattern-score × injected allocation; not oil gain prediction',
                  'Calculated as supported oil / water injection (not 0–1 display normalized)',
                  'Reconstructed with explicit mapping' if not comp['scatter_pattern'].empty else 'Unavailable: mapping required']})
        meta.to_excel(writer,sheet_name='Interpretation',index=False)
        for sheet in writer.sheets.values():
            sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
            from openpyxl.styles import Font,PatternFill
            for cell in sheet[1]:
                cell.fill=PatternFill('solid',fgColor='173553');cell.font=Font(bold=True,color='FFFFFF')
            from openpyxl.utils import get_column_letter
            for col in sheet.columns:
                maxlen=max(len(str(x.value or '')) for x in list(col)[:100])
                sheet.column_dimensions[get_column_letter(col[0].column)].width=min(40,max(13,maxlen+2))
    left,right=st.columns(2)
    with left:
        st.download_button('⬇ Export Original vs Optimized Excel',excel.getvalue(),
            file_name='Rumaila_WI_Before_After_Analysis.xlsx',
            mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            key='cmp_excel_dl',use_container_width=True,type='primary')
    with right:
        st.download_button('⬇ Export interactive plots (HTML)',_download_html(charts),
            file_name='Rumaila_WI_Optimization_Interactive_Plots.html',
            mime='text/html',key='cmp_html_dl',use_container_width=True)
    return comp
